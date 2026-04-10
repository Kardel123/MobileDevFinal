"""Generate Campus TaskHub course presentation (PowerPoint). Run: python tools/build_presentation.py"""

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
OUT = Path(__file__).resolve().parents[1] / "presentation" / "Campus_TaskHub_Presentation.pptx"

PRIMARY = RGBColor(0x15, 0x65, 0xC0)  # app blue


def add_title_slide(prs, title: str, subtitle: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    st = slide.placeholders[1]
    st.text = subtitle


def add_bullet_slide(prs, title: str, bullets: list[str]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    body = slide.shapes.placeholders[1].text_frame
    body.clear()
    for i, line in enumerate(bullets):
        if i == 0:
            p = body.paragraphs[0]
        else:
            p = body.add_paragraph()
        p.text = line
        p.level = 0
        p.font.size = Pt(18)
        if line.startswith("•") or line.startswith("–"):
            p.level = 0


def add_two_column_bullet(prs, title: str, left_title: str, left: list[str], right_title: str, right: list[str]) -> None:
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    # Title
    tx = slide.shapes.add_textbox(Inches(0.5), Inches(0.35), Inches(9), Inches(0.8))
    tx.text_frame.text = title
    tx.text_frame.paragraphs[0].font.size = Pt(32)
    tx.text_frame.paragraphs[0].font.bold = True
    tx.text_frame.paragraphs[0].font.color.rgb = PRIMARY

    def box(x, y, w, h, header, items):
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = shape.text_frame
        tf.text = header
        tf.paragraphs[0].font.size = Pt(20)
        tf.paragraphs[0].font.bold = True
        for item in items:
            p = tf.add_paragraph()
            p.text = item
            p.level = 0
            p.font.size = Pt(16)
            p.space_after = Pt(6)

    box(0.5, 1.2, 4.5, 5.5, left_title, left)
    box(5.2, 1.2, 4.5, 5.5, right_title, right)


def main() -> None:
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    add_title_slide(
        prs,
        "Campus TaskHub",
        "Mobile computing project — Flutter · Supabase · MVVM\nPresentation outline (≈7 min + demo + Q&A)",
    )

    add_bullet_slide(
        prs,
        "Logistics (reference)",
        [
            "• Target: ~12 minutes total including setup",
            "• Structure: ~7 min content + live demo + future plans, then Q&A",
            "• Tip: rehearse demo path; keep backup screenshots",
        ],
    )

    add_bullet_slide(
        prs,
        "1. Title & context",
        [
            "• Project: Campus TaskHub",
            "• Team: [Add names]",
            "• Roles: [e.g. Backend/Supabase · App & MVVM · UI/flows · QA & demo]",
        ],
    )

    add_bullet_slide(
        prs,
        "2. Project overview",
        [
            "• What: Student mobile app for tasks, schedule, and group collaboration",
            "• Problem: Scattered deadlines; need shared group work and college/subject context",
            "• Users: Enrolled students (catalog-driven registration)",
        ],
    )

    add_two_column_bullet(
        prs,
        "3. Project objectives",
        "Problems & goals",
        [
            "Centralized tasks & deadlines",
            "Groups via join codes & roles",
            "Insights on workload",
        ],
        "Technical & user-centric",
        [
            "Supabase + RLS + Auth",
            "MVVM + Provider",
            "Usable task/calendar/team flows",
        ],
    )

    add_bullet_slide(
        prs,
        "4. Core functionalities",
        [
            "• Auth (Supabase) · Onboarding: college + subject catalog + profile",
            "• Tasks: CRUD, filters, pin, progress %, status",
            "• Groups / Team: join codes, members, active group",
            "• Calendar: week grid & day agenda; classes + due dates",
            "• Academics: Tasks, Projects, Insights, Team",
            "• Dashboard & profile · Local deadline notifications (where supported)",
        ],
    )

    add_bullet_slide(
        prs,
        "5. Architecture overview",
        [
            "• Pattern: MVVM — Views observe ViewModels; Services call Supabase",
            "• State: Provider (ChangeNotifier, MultiProvider)",
            "• Backend: Supabase — Postgres, Row Level Security, Auth, RPC (e.g. join by code)",
            "• Client: Flutter (Material), supabase_flutter",
            "• See next slide for MVVM flow",
        ],
    )

    add_bullet_slide(
        prs,
        "MVVM & data flow (conceptual)",
        [
            "View (widgets) → watches → ViewModel (ChangeNotifier)",
            "ViewModel → Service (TaskService, GroupService, StudentProfileService, …)",
            "Service → Supabase client → Auth + Postgres (RLS)",
            "Models: e.g. AcademicTask, StudentAppContext — held in ViewModels",
        ],
    )

    add_bullet_slide(
        prs,
        "6. Development phases",
        [
            "• Insert your Gantt chart on this slide or the next",
            "• Suggested phases: requirements → schema & auth → MVVM shell → tasks & groups",
            "• → calendar & registration catalog → insights / team / notifications → test & polish",
        ],
    )

    add_bullet_slide(
        prs,
        "7. Live demo (rehearse)",
        [
            "• Login / signup (if time)",
            "• Registration or profile: college + subjects",
            "• Tasks: add, pin, progress, done, filters",
            "• Team: join code / groups (if stable)",
            "• Insights · Calendar (week vs day)",
            "• Optional: show error handling (e.g. invalid code / offline)",
        ],
    )

    add_bullet_slide(
        prs,
        "8. Future plans",
        [
            "• Push notifications (FCM) beyond local reminders",
            "• Localization (e.g. Filipino / English)",
            "• Deeper faculty/admin workflows · Offline or sync status",
            "• Analytics / A/B testing (optional)",
        ],
    )

    add_bullet_slide(
        prs,
        "9. Q&A (≈5 min)",
        [
            "• Why Supabase? Postgres + Auth + RLS; fast iteration for student projects",
            "• Why MVVM + Provider? Clear separation; fits Flutter patterns",
            "• Security: RLS; sensitive ops via RPC where needed",
            "• Be ready: hardest bug, trade-offs, what you’d redo",
        ],
    )

    add_title_slide(
        prs,
        "Thank you",
        "Campus TaskHub — questions?",
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"Saved: {OUT}")


if __name__ == "__main__":
    main()
